"""Export showtimes to CSV / Excel."""
from __future__ import annotations

import csv
import logging
from typing import List

from .models import MovieShowtimes, Showtime

log = logging.getLogger(__name__)


def _rows(all_ms: List[MovieShowtimes]) -> List[dict]:
    rows = []
    for ms in all_ms:
        for st in ms.showtimes:
            rows.append(
                {
                    "Date": st.date,
                    "Movie": ms.movie.title,
                    "Slug": ms.movie.slug,
                    "Cinema": st.cinema,
                    "Format": st.format_name,
                    "Time": st.time,
                    "Status": "Available" if st.available else "Sold out",
                    "Total Seats": st.total_seats,
                    "Available Seats": st.available_seats,
                    "Sold Seats": st.sold_seats,
                    "Booking URL": st.booking_url,
                }
            )
    return rows


def export_csv(path: str, all_ms: List[MovieShowtimes]) -> None:
    rows = _rows(all_ms)
    if not rows:
        log.warning("No rows to export")
        return
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    log.info("Exported %d rows to %s", len(rows), path)


def export_excel(path: str, all_ms: List[MovieShowtimes]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl not installed; use CSV export instead")
    rows = _rows(all_ms)
    wb = Workbook()
    ws = wb.active
    ws.title = "Showtimes"
    if not rows:
        wb.save(path)
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    # header styling
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r[h] for h in headers])
    # auto width
    for c, h in enumerate(headers, start=1):
        maxlen = max(len(str(h)), max((len(str(r[h])) for r in rows), default=0))
        ws.column_dimensions[get_column_letter(c)].width = min(maxlen + 2, 45)
    # color status
    status_idx = headers.index("Status") + 1
    for row in ws.iter_rows(min_row=2):
        cell = row[status_idx - 1]
        if cell.value == "Available":
            cell.fill = PatternFill("solid", fgColor="C6EFCE")
        elif cell.value == "Sold out":
            cell.fill = PatternFill("solid", fgColor="FFC7CE")
    wb.save(path)
    log.info("Exported %d rows to %s", len(rows), path)
